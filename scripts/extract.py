"""Extract plain text from a document.

Called by the Rust `read_document` tool for formats Rust cannot read directly.
Prints text to stdout; exits non-zero with a message on stderr if the format is
unsupported or a library is missing.

Requires: pymupdf, python-docx, openpyxl, python-pptx (optional per format).
"""
import sys
from pathlib import Path

MAX_CHARS = 200_000


def from_pdf(path):
    import fitz  # pymupdf

    with fitz.open(path) as doc:
        pages = [f"--- page {i + 1} ---\n{p.get_text()}" for i, p in enumerate(doc)]
    return "\n".join(pages)


def from_docx(path):
    import docx

    doc = docx.Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    # Tables carry most of the content in many documents, so include them.
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def from_xlsx(path):
    import openpyxl

    book = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    out = []
    for sheet in book.worksheets:
        out.append(f"--- sheet: {sheet.title} ---")
        for row in sheet.iter_rows(values_only=True):
            if any(c is not None for c in row):
                out.append("\t".join("" if c is None else str(c) for c in row))
    return "\n".join(out)


def from_pptx(path):
    from pptx import Presentation

    deck = Presentation(str(path))
    out = []
    for i, slide in enumerate(deck.slides, 1):
        out.append(f"--- slide {i} ---")
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                out.append(shape.text_frame.text)
    return "\n".join(out)


EXTRACTORS = {
    ".pdf": from_pdf,
    ".docx": from_docx,
    ".doc": from_docx,
    ".xlsx": from_xlsx,
    ".xls": from_xlsx,
    ".pptx": from_pptx,
}


def main():
    if len(sys.argv) < 2:
        print("usage: extract.py <file>", file=sys.stderr)
        return 2

    path = Path(sys.argv[1])
    if not path.exists():
        print(f"no such file: {path}", file=sys.stderr)
        return 1

    extractor = EXTRACTORS.get(path.suffix.lower())
    if extractor is None:
        print(f"no extractor for {path.suffix} files", file=sys.stderr)
        return 1

    try:
        text = extractor(path)
    except ImportError as exc:
        print(f"missing library for {path.suffix}: {exc}", file=sys.stderr)
        return 1
    except Exception as exc:  # noqa: BLE001 - report and let the agent adapt
        print(f"could not read {path.name}: {exc}", file=sys.stderr)
        return 1

    sys.stdout.write(text[:MAX_CHARS])
    return 0


if __name__ == "__main__":
    sys.exit(main())
